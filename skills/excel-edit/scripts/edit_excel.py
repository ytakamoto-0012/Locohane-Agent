"""ops（操作）のリストを適用して xlsx/xlsm を新規作成または編集する。

excel-edit スキルの実行スクリプト（progressive disclosure 第3段階）。
run_script から
    python edit_excel.py <path> --ops-json "<JSON配列>" [--new] [--overwrite] [--output <別名保存先>]
または
    python edit_excel.py <path> --ops-file <JSONファイルのパス> [--new] [--overwrite] [--output ...]
の形で呼ばれる（このプロジェクトには汎用のファイル書き込みツールが無いため、
LLMが組み立てたJSONをそのまま --ops-json 引数として渡せるようにしている）。

--new を付けると新規ワークブックを作成する（既存ファイルがある場合は
--overwrite 必須）。--new を付けない場合は <path> を読み込んで編集する
（存在しなければエラー）。--output を省略した場合は <path> へ上書き保存する。

ops の各要素の形式や対応opの一覧は SKILL.md を参照。実装（各opの処理）は
scripts/_ops.py の OP_HANDLERS を参照。
"""

from __future__ import annotations

import argparse
import json
import sys
import traceback
from pathlib import Path

# office_shared/excel_common.py から共有ヘルパーを import する（1-B 相互import方式）
_OFFICE_SHARED = Path(__file__).resolve().parent.parent.parent / "office_shared"
if str(_OFFICE_SHARED) not in sys.path:
    sys.path.append(str(_OFFICE_SHARED))
from excel_common import backup_before_overwrite, read_text_with_fallback, register_output_path, setup_utf8_stdio  # noqa: E402
from _ops import apply_op  # noqa: E402


def _load_json_arg(args: argparse.Namespace) -> str:
    if args.ops_json is not None:
        return args.ops_json
    data_path = Path(args.ops_file)
    if not data_path.is_file():
        raise FileNotFoundError(f"ops file not found: {args.ops_file}")
    return read_text_with_fallback(data_path)


def main() -> int:
    setup_utf8_stdio()
    parser = argparse.ArgumentParser()
    parser.add_argument("path", help="xlsx/xlsm path to edit (the path to create with --new)")
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--ops-json", help="pass the JSON array of operations to apply as is")
    group.add_argument("--ops-file", help="path of a file with the JSON array of operations to apply")
    parser.add_argument("--new", action="store_true", help="create a new workbook without reading an existing file")
    parser.add_argument("--overwrite", action="store_true", help="with --new, allow overwriting if the output already exists")
    parser.add_argument("--output", default=None, help="output path (if omitted, path is overwritten)")
    args = parser.parse_args()

    path = Path(args.path)
    ext = path.suffix.lower()
    if ext not in (".xlsx", ".xlsm"):
        print(f"Only the .xlsx/.xlsm extensions are supported: {ext}", file=sys.stderr)
        return 1

    output_path = Path(args.output) if args.output else path

    try:
        raw = _load_json_arg(args)
    except FileNotFoundError as e:
        print(str(e), file=sys.stderr)
        return 1

    try:
        ops = json.loads(raw)
    except json.JSONDecodeError as e:
        print(f"Failed to parse the ops JSON: {e}", file=sys.stderr)
        return 1
    if not isinstance(ops, list):
        print("ops must be a JSON array", file=sys.stderr)
        return 1

    import openpyxl

    if args.new:
        if path.exists() and not args.overwrite:
            print(f"The output already exists (specify --overwrite to overwrite it): {args.path}", file=sys.stderr)
            return 1
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
    else:
        if not path.is_file():
            print(f"File to edit not found (specify --new to create a new one): {args.path}", file=sys.stderr)
            return 1
        try:
            wb = openpyxl.load_workbook(str(path), keep_vba=(ext == ".xlsm"))
        except Exception as e:  # noqa: BLE001 - 破損ファイル等、openpyxlが送出する多様な例外を丸めて報告する
            print(f"Failed to read the file: {e}\n{traceback.format_exc()}", file=sys.stderr)
            return 1

    for idx, op in enumerate(ops):
        if not isinstance(op, dict) or "op" not in op:
            print(f"ops[{idx}] is invalid (it must be an object with an 'op' key): {op!r}", file=sys.stderr)
            return 1
        try:
            apply_op(wb, op)
        except KeyError as e:
            print(f"ops[{idx}] (op={op.get('op')!r}) is missing the required key {e}", file=sys.stderr)
            return 1
        except (ValueError, TypeError) as e:
            print(f"Failed to apply ops[{idx}] (op={op.get('op')!r}): {e}", file=sys.stderr)
            return 1

    if not wb.sheetnames:
        print("In the end there are no sheets at all (include add_sheet or review --new)", file=sys.stderr)
        return 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    backup_path = backup_before_overwrite(output_path)
    try:
        wb.save(str(output_path))
    except OSError as e:
        print(f"Failed to save the file: {e}", file=sys.stderr)
        return 1

    result = {
        "path": str(output_path),
        "backup_path": str(backup_path) if backup_path else None,
        "sheets": wb.sheetnames,
        "applied_ops": len(ops),
    }
    path_memory = register_output_path(output_path, description="file generated/updated by edit_excel")
    if path_memory:
        result["path_memory"] = path_memory
    print(json.dumps(result, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
