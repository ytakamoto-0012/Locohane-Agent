"""英語版 Windows など、日本語以外の環境でも落ちないことのテスト。"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest

from src.config import PROJECT_ROOT
from src.tools._file_tools_common import read_text_with_fallback

PDF_SCRIPTS = PROJECT_ROOT / "skills" / "pdf-tools" / "scripts"


@pytest.fixture
def create_pdf(monkeypatch):
    monkeypatch.syspath_prepend(str(PDF_SCRIPTS))
    import create_pdf as module

    yield module
    # pdf-tools の _common を残すと、同名の skill-creator/scripts/_common を import するテストが失敗する
    for name in ("create_pdf", "_common"):
        sys.modules.pop(name, None)


def test_pdf_uses_latin_fonts_when_japanese_fonts_are_missing(create_pdf, monkeypatch) -> None:
    if not (create_pdf._WINDOWS_FONTS_DIR / "arial.ttf").exists():
        pytest.skip("Windows fonts are not available")
    monkeypatch.setattr(create_pdf, "_JP_GOTHIC_CANDIDATES", ["missing-gothic.ttc"])
    monkeypatch.setattr(create_pdf, "_JP_MINCHO_CANDIDATES", ["missing-mincho.ttc"])

    # 英語だけの文書は欧文フォントで作れる
    create_pdf.register_jp_fonts("<h1>Report</h1><p>Hello</p>")

    # 日本語を含む文書は、豆腐化けした PDF を返さないよう明示的にエラーにする
    with pytest.raises(RuntimeError, match="Japanese"):
        create_pdf.register_jp_fonts("<p>" + chr(0x65E5) + chr(0x672C) + "</p>")


def test_text_fallback_reads_cp932_and_normalizes_newlines(tmp_path: Path) -> None:
    sjis = tmp_path / "sjis.txt"
    sjis.write_bytes(bytes([0x93, 0xFA, 0x96, 0x7B]) + b"\r\nx")
    assert read_text_with_fallback(sjis) == chr(0x65E5) + chr(0x672C) + "\nx"


def test_text_fallback_uses_ansi_code_page(tmp_path: Path, monkeypatch) -> None:
    # 英語版 Windows（ANSI コードページ cp1252）の Latin-1 系ファイルを読める
    monkeypatch.setattr("locale.getencoding", lambda: "cp1252")
    latin = tmp_path / "latin.txt"
    latin.write_bytes(b"caf\xe9")
    assert read_text_with_fallback(latin) == "caf" + chr(0xE9)


@pytest.mark.parametrize("text", ["Gr" + chr(0xF6) + chr(0xDF) + "e Stra" + chr(0xDF) + "e", "Se" + chr(0xF1) + "or a" + chr(0xF1) + "o"])
def test_text_fallback_prefers_ansi_code_page_over_cp932(tmp_path: Path, monkeypatch, text: str) -> None:
    # cp1252 の文書の多くは cp932 としてもエラーなく読めてしまう（"Größe" → "Gre"）。ANSI を先に試す
    monkeypatch.setattr("locale.getencoding", lambda: "cp1252")
    latin = tmp_path / "latin.txt"
    latin.write_bytes(text.encode("cp1252"))
    assert read_text_with_fallback(latin) == text

    from src.vector_db.extract import _decode_text

    assert _decode_text(text.encode("cp1252")) == text


def test_text_fallback_still_reads_shift_jis_on_english_windows(tmp_path: Path, monkeypatch) -> None:
    # 句読点（先頭バイト 0x81）を含む Shift-JIS は cp1252 で読めないため、cp932 へ進んで正しく読める
    monkeypatch.setattr("locale.getencoding", lambda: "cp1252")
    text = chr(0x65E5) + chr(0x672C) + chr(0x3002)  # 日本。
    sjis = tmp_path / "sjis.txt"
    sjis.write_bytes(text.encode("cp932"))
    assert read_text_with_fallback(sjis) == text


@pytest.mark.parametrize("skill", ["docx-create", "docx-edit"])
def test_docx_default_fonts_use_latin_fonts_without_japanese(monkeypatch, skill: str) -> None:
    monkeypatch.syspath_prepend(str(PROJECT_ROOT / "skills" / skill / "scripts"))
    sys.modules.pop("_blocks", None)
    import _blocks
    from docx import Document
    from docx.oxml.ns import qn

    try:
        def fonts(text: str) -> tuple[str | None, str | None]:
            _blocks.use_latin_fonts_unless_japanese(text)
            run = Document().add_paragraph().add_run("x")
            _blocks._set_east_asian_font(run.font, run._element, _blocks.HEADING_FONT)
            return run.font.name, run._element.rPr.rFonts.get(qn("w:eastAsia"))

        # 英語だけの文書: 欧文側は Arial、東アジア側は游ゴシックのまま
        assert fonts('{"text": "Report"}') == ("Arial", _blocks.HEADING_FONT)
        # 日本語を含む文書: 従来どおり欧文側も游ゴシック
        assert fonts('{"text": "' + chr(0x5831) + chr(0x544A) + '"}') == (_blocks.HEADING_FONT, _blocks.HEADING_FONT)
    finally:
        sys.modules.pop("_blocks", None)
