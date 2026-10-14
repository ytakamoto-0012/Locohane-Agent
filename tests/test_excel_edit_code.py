"""excel-editスキル（skills/excel-edit/scripts/edit_excel.py）のopenpyxlコード実行方式。

ops JSON方式からの比較検証用に、LLMが書いたopenpyxlコードを --code/--code-file で
受け取って実行する方式へ切り替えた。読み込み・保存・バックアップはスクリプト側が
担い、コードは渡された wb を書き換えるだけにする契約を検証する。
"""

import importlib
import json
import sys
from pathlib import Path

import openpyxl
import pytest

_EXCEL_EDIT_SCRIPTS = Path(__file__).resolve().parent.parent / "skills" / "excel-edit" / "scripts"
_OFFICE_SHARED = Path(__file__).resolve().parent.parent / "skills" / "office_shared"


@pytest.fixture
def edit_excel():
    # docx-edit にも同名の _ops.py があるため、sys.modules のキャッシュを捨ててから読み直す
    for name in ("edit_excel", "xlsx_helpers", "_ops", "_style"):
        sys.modules.pop(name, None)
    for d in (_EXCEL_EDIT_SCRIPTS, _OFFICE_SHARED):
        if str(d) in sys.path:
            sys.path.remove(str(d))
        sys.path.insert(0, str(d))
    return importlib.import_module("edit_excel")


def _run(edit_excel, monkeypatch, *argv):
    monkeypatch.setattr(sys, "argv", ["edit_excel.py", *map(str, argv)])
    return edit_excel.main()


def test_new_book_with_helpers_is_saved(edit_excel, tmp_path, monkeypatch, capsys):
    target = tmp_path / "book.xlsx"
    code = """
ws = wb.create_sheet("売上")
print(write_table(ws, "A1", [["月","売上"], ["4月",120], ["5月",135]], theme="navy"))
"""
    assert _run(edit_excel, monkeypatch, target, "--new", "--code", code) == 0

    result = json.loads(capsys.readouterr().out.strip().splitlines()[-1])
    assert result["sheets"] == ["売上"]
    assert result["code_output"].strip() == "A1:B3"
    ws = openpyxl.load_workbook(target)["売上"]
    assert ws["B3"].value == 135
    assert ws["A1"].fill.fgColor.rgb.endswith("1E2761")
    assert ws.freeze_panes == "A2"


def test_code_file_is_accepted(edit_excel, tmp_path, monkeypatch):
    target = tmp_path / "book.xlsx"
    code_file = tmp_path / "edit.py"
    code_file.write_text('ws = wb.create_sheet("S")\nws["A1"] = "値"\n', encoding="utf-8")

    assert _run(edit_excel, monkeypatch, target, "--new", "--code-file", code_file) == 0
    assert openpyxl.load_workbook(target)["S"]["A1"].value == "値"


def test_failing_code_reports_line_and_saves_nothing(edit_excel, tmp_path, monkeypatch, capsys):
    target = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = "元の値"
    wb.save(target)
    before = target.read_bytes()

    code = 'ws = wb.active\nws["A1"] = "書き換え"\nws["B1"] = 1 / 0\n'
    assert _run(edit_excel, monkeypatch, target, "--code", code) == 1

    err = capsys.readouterr().err
    assert "ZeroDivisionError" in err
    assert "コード3行目" in err
    assert "保存していません" in err
    assert target.read_bytes() == before
    assert not list(tmp_path.glob("*.bak_*"))


def test_unknown_sheet_lists_sheetnames(edit_excel, tmp_path, monkeypatch, capsys):
    target = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "集計"
    wb.save(target)

    assert _run(edit_excel, monkeypatch, target, "--code", 'wb["売上"]["A1"] = 1') == 1
    assert "['集計']" in capsys.readouterr().err


def test_syntax_error_reports_line(edit_excel, tmp_path, monkeypatch, capsys):
    target = tmp_path / "book.xlsx"
    assert _run(edit_excel, monkeypatch, target, "--new", "--code", 'ws = wb.create_sheet("S")\nfor i in range(3)\n    pass') == 1
    err = capsys.readouterr().err
    assert "構文エラー" in err
    assert "2行目" in err


def test_save_inside_code_is_forbidden(edit_excel, tmp_path, monkeypatch, capsys):
    target = tmp_path / "book.xlsx"
    other = tmp_path / "other.xlsx"
    code = f'wb.create_sheet("S")\nwb.save(r"{other}")'
    assert _run(edit_excel, monkeypatch, target, "--new", "--code", code) == 1
    assert "wb.save()" in capsys.readouterr().err
    assert not other.exists()
    assert not target.exists()


def test_replacing_wb_is_rejected(edit_excel, tmp_path, monkeypatch, capsys):
    target = tmp_path / "book.xlsx"
    assert _run(edit_excel, monkeypatch, target, "--new", "--code", "wb = openpyxl.Workbook()") == 1
    assert "差し替えないで" in capsys.readouterr().err
    assert not target.exists()


def test_new_book_without_sheets_is_rejected(edit_excel, tmp_path, monkeypatch, capsys):
    target = tmp_path / "book.xlsx"
    assert _run(edit_excel, monkeypatch, target, "--new", "--code", "x = 1") == 1
    assert "シートが1つも存在しません" in capsys.readouterr().err


def test_new_requires_overwrite_for_existing_file(edit_excel, tmp_path, monkeypatch, capsys):
    target = tmp_path / "book.xlsx"
    openpyxl.Workbook().save(target)
    assert _run(edit_excel, monkeypatch, target, "--new", "--code", 'wb.create_sheet("S")') == 1
    assert "--overwrite" in capsys.readouterr().err


def test_code_execution_disabled_by_config(edit_excel, tmp_path, monkeypatch, capsys):
    monkeypatch.setenv("AGENT_CODE_EXECUTION_ENABLED", "0")
    target = tmp_path / "book.xlsx"
    assert _run(edit_excel, monkeypatch, target, "--new", "--code", 'wb.create_sheet("S")') == 1
    assert "code_execution_enabled" in capsys.readouterr().err
    assert not target.exists()


def test_list_groups_merges_repeated_values_and_merged_cells(edit_excel):
    import xlsx_helpers

    wb = openpyxl.Workbook()
    ws = wb.active
    for row in [["月"], ["5月"], ["7月"], ["7月"], ["8月"], [None], ["9月"]]:
        ws.append(row)
    ws.merge_cells("A5:A6")

    groups = [(g["value"], g["start_row"], g["end_row"]) for g in xlsx_helpers.list_groups(ws, "A")]

    assert groups == [("月", 1, 1), ("5月", 2, 2), ("7月", 3, 4), ("8月", 5, 6), ("9月", 7, 7)]
    assert xlsx_helpers.find_group(ws, "A", "7月") == (3, 4)
    with pytest.raises(ValueError, match="見つかりません"):
        xlsx_helpers.find_group(ws, "A", "10月")


def test_insert_rows_helper_keeps_merged_cells(edit_excel):
    import xlsx_helpers

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.merge_cells("A3:A4")

    xlsx_helpers.insert_rows(ws, 2, 2)

    assert [str(r) for r in ws.merged_cells.ranges] == ["A5:A6"]


def test_write_table_row_styles_are_not_overwritten_by_banding(edit_excel):
    import xlsx_helpers

    wb = openpyxl.Workbook()
    ws = wb.active
    rows = [["月", "行事"], ["1月", "a"], ["1月", "b"], ["2月", "c"]]
    styles = [None, {"fill_color": "DAEEF3"}, {"fill_color": "DAEEF3"}, {"fill_color": "F8CBAD"}]

    xlsx_helpers.write_table(ws, "A1", rows, row_styles=styles)

    assert [ws.cell(row=r, column=1).fill.fgColor.rgb[-6:] for r in (2, 3, 4)] == ["DAEEF3", "DAEEF3", "F8CBAD"]
