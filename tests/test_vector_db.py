"""src/vector_db（抽出・チャンク分割・埋め込み・保存先・取り込み・検索）のテスト。"""

from __future__ import annotations

import os

import httpx
import pytest

from _vector_db_fakes import DIM, FakeEmbeddingServer, FakeQdrant, fake_vector, make_config
from src.config import LLMEndpoint
from src.vector_db import catalog as vcat
from src.vector_db import embeddings, ingest, search
from src.vector_db.chunking import split_text
from src.vector_db.embeddings import EmbeddingError, EmbeddingStopped, embed_images, embed_texts, open_client, shared_client
from src.vector_db.extract import ExtractError, NotTextError, extract_segments
from src.vector_db.qdrant_store import QDRANT_PRESETS, QdrantStore
from src.vector_db.sqlite_vec_store import SqliteVecStore
from src.vector_db.store import ChunkRecord, StoreError


# --- チャンク分割 ---------------------------------------------------------------


def test_split_text_keeps_short_paragraphs_together():
    assert split_text("段落1\n\n段落2", chunk_size=100, overlap=10) == ["段落1\n段落2"]


def test_split_text_respects_size_and_overlap():
    text = "\n\n".join(f"{i:02d}" + "あ" * 40 for i in range(10))
    chunks = split_text(text, chunk_size=100, overlap=20)
    assert len(chunks) > 1
    assert all(len(c) <= 100 for c in chunks)
    # 2つ目以降は直前のチャンクの末尾を重ねている
    assert chunks[1].startswith(chunks[0][-20:])


def test_split_text_splits_long_line_by_characters():
    chunks = split_text("x" * 250, chunk_size=100, overlap=0)
    assert [len(c) for c in chunks] == [100, 100, 50]


def test_split_text_rejects_overlap_not_smaller_than_size():
    with pytest.raises(ValueError):
        split_text("abc", chunk_size=10, overlap=10)


# --- テキスト抽出 ---------------------------------------------------------------


def test_extract_text_falls_back_to_cp932(tmp_path):
    path = tmp_path / "a.txt"
    path.write_bytes("溶接の基準".encode("cp932"))
    assert extract_segments(path)[0].text == "溶接の基準"


def test_extract_office_and_pdf_formats(tmp_path):
    import docx
    import openpyxl
    from pptx import Presentation
    from reportlab.pdfgen import canvas

    d = docx.Document()
    d.add_paragraph("docx body")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "cell1"
    d.save(tmp_path / "a.docx")
    assert "docx body" in extract_segments(tmp_path / "a.docx")[0].text
    assert "cell1" in extract_segments(tmp_path / "a.docx")[0].text

    wb = openpyxl.Workbook()
    wb.active.title = "集計"
    wb.active.append(["品番", 10])
    wb.create_sheet("空")
    wb.save(tmp_path / "a.xlsx")
    segments = extract_segments(tmp_path / "a.xlsx")
    assert [(s.locator, s.text) for s in segments] == [("Sheet: 集計", "品番\t10")]

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "slide title"
    slide.notes_slide.notes_text_frame.text = "speaker note"
    prs.save(tmp_path / "a.pptx")
    segments = extract_segments(tmp_path / "a.pptx")
    assert segments[0].locator == "Slide 1"
    assert "slide title" in segments[0].text and "speaker note" in segments[0].text

    c = canvas.Canvas(str(tmp_path / "a.pdf"))
    c.drawString(72, 720, "page one text")
    c.showPage()
    c.drawString(72, 720, "page two text")
    c.save()
    segments = extract_segments(tmp_path / "a.pdf")
    assert [s.locator for s in segments] == ["p.1", "p.2"]
    assert "page two text" in segments[1].text


_XDR = 'xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
_MC = 'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'


def _shape(text: str) -> str:
    return f"<xdr:sp><xdr:txBody><a:p><a:r><a:t>{text}</a:t></a:r></a:p></xdr:txBody></xdr:sp>"


def _cell_anchor(row: int, col: int, inner: str) -> str:
    return f"<xdr:twoCellAnchor><xdr:from><xdr:col>{col}</xdr:col><xdr:row>{row}</xdr:row></xdr:from>{inner}</xdr:twoCellAnchor>"


def _add_parts(path, parts: dict[str, str]) -> None:
    """xlsx（ZIP）にパーツを足す（openpyxl では書けない図形・スレッド形式のコメントのテスト用）。"""
    import zipfile

    with zipfile.ZipFile(path, "a") as z:
        for name, xml in parts.items():
            z.writestr(name, xml)


def _rels(*items: tuple[str, str, str]) -> str:
    body = "".join(
        f'<Relationship Id="{rel_id}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/{kind}" Target="{target}"/>'
        for rel_id, kind, target in items
    )
    return f'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">{body}</Relationships>'


def test_extract_xlsx_inserts_shapes_and_comments_at_their_rows(tmp_path):
    import openpyxl
    from openpyxl.comments import Comment

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "集計"
    for row in (["品番", 10], ["型式", 20], ["数量", 30]):
        ws.append(row)
    ws["B2"].comment = Comment("前年比で確認", "yt")
    path = tmp_path / "a.xlsx"
    wb.save(path)
    drawing = (
        f"<xdr:wsDr {_XDR} {_MC}>"
        + _cell_anchor(0, 2, _shape("承認欄"))  # C1
        # グループの中の図形。新旧両形式（mc:Fallback）を持つ図形は1回だけ読む。
        + _cell_anchor(1, 0, f"<xdr:grpSp>{_shape('段取り手順')}<mc:AlternateContent><mc:Choice>{_shape('注意')}</mc:Choice>"
                             f"<mc:Fallback>{_shape('注意')}</mc:Fallback></mc:AlternateContent></xdr:grpSp>")  # A2
        + _cell_anchor(9, 0, _shape("欄外のメモ"))  # A10（データより下）
        + _cell_anchor(2, 0, '<xdr:graphicFrame><a:graphic><a:graphicData><c:chart xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
                             'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rId1"/></a:graphicData></a:graphic></xdr:graphicFrame>')
        + f"<xdr:absoluteAnchor>{_shape('ロゴ')}</xdr:absoluteAnchor>"
        + "</xdr:wsDr>"
    )
    chart = (
        '<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        "<c:chart><c:title><c:tx><c:rich><a:p><a:r><a:t>月別の出荷数</a:t></a:r></a:p></c:rich></c:tx></c:title></c:chart></c:chartSpace>"
    )
    import zipfile

    with zipfile.ZipFile(path) as z:
        sheet_rels = z.read("xl/worksheets/_rels/sheet1.xml.rels").decode()
    sheet_rels = sheet_rels.replace("</Relationships>", _rels(("rIdD", "drawing", "../drawings/drawing9.xml")).split(">", 1)[1])
    # 既存の sheet1.xml.rels（コメント）を差し替えるため、新しい ZIP に書き直す。
    rebuilt = tmp_path / "b.xlsx"
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(rebuilt, "w") as dst:
        for item in src.infolist():
            data = sheet_rels.encode() if item.filename == "xl/worksheets/_rels/sheet1.xml.rels" else src.read(item.filename)
            dst.writestr(item, data)
    _add_parts(rebuilt, {
        "xl/drawings/drawing9.xml": drawing,
        "xl/drawings/_rels/drawing9.xml.rels": _rels(("rId1", "chart", "../charts/chart1.xml")),
        "xl/charts/chart1.xml": chart,
    })
    [segment] = extract_segments(rebuilt)
    # 図形・コメントはその行の後ろ（同じ行なら列の順）、データより下の図形は末尾、位置で置いた図形は最後。
    assert segment.text.split("\n") == [
        "品番\t10",
        "[Shape C1] 承認欄",
        "型式\t20",
        "[Shape A2] 段取り手順",
        "[Shape A2] 注意",
        "[Comment B2] 前年比で確認",
        "数量\t30",
        "[Shape A3] 月別の出荷数",
        "[Shape A10] 欄外のメモ",
        "[Shape] ロゴ",
    ]


def test_extract_xlsx_prefers_threaded_comments_over_their_legacy_copies(tmp_path):
    import zipfile

    import openpyxl
    from openpyxl.comments import Comment

    wb = openpyxl.Workbook()
    wb.active.append(["品番"])
    # スレッド形式のコメントを付けると、Excel は古いバージョン向けの説明文入りの複製も古い形式で書く。
    wb.active["A1"].comment = Comment("[Threaded comment] Your version of Excel allows you to read this threaded comment", "tc")
    path = tmp_path / "a.xlsx"
    wb.save(path)
    rebuilt = tmp_path / "b.xlsx"
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(rebuilt, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/_rels/sheet1.xml.rels":
                data = data.replace(b"</Relationships>", _rels(("rIdT", "threadedComment", "../threadedComments/tc1.xml")).split(">", 1)[1].encode())
            dst.writestr(item, data)
    _add_parts(rebuilt, {
        "xl/threadedComments/tc1.xml": '<ThreadedComments xmlns="http://schemas.microsoft.com/office/spreadsheetml/2018/threadedcomments">'
        '<threadedComment ref="A1" id="{1}"><text>単価は税抜</text></threadedComment></ThreadedComments>'
    })
    assert extract_segments(rebuilt)[0].text == "品番\n[Comment A1] 単価は税抜"


def test_extract_docx_reads_text_boxes_headers_footers_and_comments(tmp_path):
    import docx
    from docx.oxml import parse_xml

    d = docx.Document()
    para = d.add_paragraph("docx body")
    d.sections[0].header.paragraphs[0].text = "社外秘"
    d.sections[0].footer.paragraphs[0].text = "品質保証部"
    d.add_comment(para.runs, text="要確認", author="yt")
    # 新旧両形式（mc:Fallback に VML の複製）を持つテキストボックス。文字は1回だけ読む。
    textbox = (
        '<w:txbxContent xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:p><w:r><w:t>{0}</w:t></w:r></w:p></w:txbxContent>'
    )
    d.element.body.insert(
        0,
        parse_xml(
            '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"><w:r><mc:AlternateContent>'
            f"<mc:Choice>{textbox.format('吹き出しの説明')}</mc:Choice><mc:Fallback>{textbox.format('吹き出しの説明')}</mc:Fallback>"
            "</mc:AlternateContent></w:r></w:p>"
        ),
    )
    d.save(tmp_path / "a.docx")
    lines = extract_segments(tmp_path / "a.docx")[0].text.split("\n")
    assert "docx body" in lines
    assert lines.count("[Text box] 吹き出しの説明") == 1
    assert {"[Header] 社外秘", "[Footer] 品質保証部", "[Comment] 要確認"} <= set(lines)


def test_extract_pptx_reads_text_in_nested_group_shapes(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text_frame.text = "工程A"
    inner = group.shapes.add_group_shape()
    inner.shapes.add_textbox(Inches(1), Inches(2), Inches(2), Inches(1)).text_frame.text = "工程B"
    prs.save(tmp_path / "a.pptx")
    assert extract_segments(tmp_path / "a.pptx")[0].text.split("\n") == ["工程A", "工程B"]


def test_extract_xlsx_reads_legacy_notes_on_cells_without_threaded_comments(tmp_path):
    import zipfile

    import openpyxl
    from openpyxl.comments import Comment

    wb = openpyxl.Workbook()
    wb.active.append(["品番", "", "単価"])
    # Excel 365 は、スレッド形式のコメント（A1）の互換用の複製と、従来形式の「メモ」（C1）を同じ comments1.xml に書く。
    wb.active["A1"].comment = Comment("[Threaded comment] Your version of Excel allows you to read this threaded comment", "tc={1}")
    wb.active["C1"].comment = Comment("暫定の値", "yt")
    wb.save(tmp_path / "a.xlsx")
    rels_path = "xl/worksheets/_rels/sheet1.xml.rels"
    rebuilt = tmp_path / "b.xlsx"
    with zipfile.ZipFile(tmp_path / "a.xlsx") as src, zipfile.ZipFile(rebuilt, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == rels_path:
                data = data.replace(b"</Relationships>", _rels(("rIdT", "threadedComment", "../threadedComments/tc1.xml")).split(">", 1)[1].encode())
            dst.writestr(item, data)
    _add_parts(rebuilt, {
        "xl/threadedComments/tc1.xml": '<ThreadedComments xmlns="http://schemas.microsoft.com/office/spreadsheetml/2018/threadedcomments">'
        '<threadedComment ref="A1" id="{1}"><text>単価は税抜</text></threadedComment></ThreadedComments>'
    })
    assert extract_segments(rebuilt)[0].text.split("\n") == ["品番\t\t単価", "[Comment A1] 単価は税抜", "[Comment C1] 暫定の値"]


def _xlsx_with_drawing(path, drawing: str, extra_parts: dict[str, str] | None = None):
    """3行のシートに図形の一覧（drawing1.xml）を付けた xlsx を作る。"""
    import openpyxl

    wb = openpyxl.Workbook()
    for value in ("r1", "r2", "r3"):
        wb.active.append([value])
    wb.save(path)
    _add_parts(path, {
        "xl/worksheets/_rels/sheet1.xml.rels": _rels(("rIdD", "drawing", "../drawings/drawing1.xml")),
        "xl/drawings/drawing1.xml": f"<xdr:wsDr {_XDR} {_MC}>{drawing}</xdr:wsDr>",
        **(extra_parts or {}),
    })
    return path


def test_extract_xlsx_reads_anchors_wrapped_in_alternate_content(tmp_path):
    # スライサー・フォームコントロール等は、アンカーごと mc:AlternateContent に包まれる。
    drawing = (
        f'<mc:AlternateContent><mc:Choice Requires="a14">{_cell_anchor(1, 1, _shape("包まれた図形"))}</mc:Choice>'
        f"<mc:Fallback>{_cell_anchor(1, 1, _shape('包まれた図形'))}</mc:Fallback></mc:AlternateContent>"
    )
    path = _xlsx_with_drawing(tmp_path / "a.xlsx", drawing)
    assert extract_segments(path)[0].text.split("\n") == ["r1", "r2", "[Shape B2] 包まれた図形", "r3"]


def test_extract_xlsx_keeps_other_shapes_when_an_anchor_is_broken(tmp_path):
    broken = f"<xdr:twoCellAnchor><xdr:from><xdr:col>0</xdr:col><xdr:row></xdr:row></xdr:from>{_shape('位置不明')}</xdr:twoCellAnchor>"
    path = _xlsx_with_drawing(tmp_path / "a.xlsx", broken + _cell_anchor(1, 0, _shape("正常な図形")))
    # セルの位置が読めない図形は、位置で置いた図形と同じくシートの末尾に付ける。
    assert extract_segments(path)[0].text.split("\n") == ["r1", "r2", "[Shape A2] 正常な図形", "r3", "[Shape] 位置不明"]


def test_extract_xlsx_reads_chart_titles_that_refer_to_cells(tmp_path):
    c = 'xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"'
    frame = (
        f'<xdr:graphicFrame><a:graphic><a:graphicData><c:chart {c} '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rId1"/></a:graphicData></a:graphic></xdr:graphicFrame>'
    )
    chart = (
        f"<c:chartSpace {c}><c:chart><c:title><c:tx><c:strRef><c:f>Sheet!$A$1</c:f><c:strCache><c:ptCount val=\"1\"/>"
        '<c:pt idx="0"><c:v>セル参照のタイトル</c:v></c:pt></c:strCache></c:strRef></c:tx></c:title></c:chart></c:chartSpace>'
    )
    path = _xlsx_with_drawing(tmp_path / "a.xlsx", _cell_anchor(0, 0, frame), {
        "xl/drawings/_rels/drawing1.xml.rels": _rels(("rId1", "chart", "../charts/chart1.xml")),
        "xl/charts/chart1.xml": chart,
    })
    assert extract_segments(path)[0].text.split("\n") == ["r1", "[Shape A1] セル参照のタイトル", "r2", "r3"]


def test_extract_xlsx_reads_chart_titles_on_chartsheets(tmp_path):
    import openpyxl
    from openpyxl.chart import BarChart, Reference

    wb = openpyxl.Workbook()
    for row in (["a", 1], ["b", 2]):
        wb.active.append(row)
    chart = BarChart()
    chart.title = "月別の出荷数"
    chart.add_data(Reference(wb.active, min_col=2, min_row=1, max_row=2))
    wb.create_chartsheet("グラフ").add_chart(chart)
    wb.create_chartsheet("空のグラフ")  # グラフを置いていないグラフシート（関係のパーツが無く、openpyxl がそのままでは開けない）
    wb.save(tmp_path / "a.xlsx")
    segments = extract_segments(tmp_path / "a.xlsx")
    assert [(s.locator, s.text) for s in segments] == [("Sheet: Sheet", "a\t1\nb\t2"), ("Sheet: グラフ", "[Shape] 月別の出荷数")]


def test_extract_xlsx_ignores_dimension_spanning_the_whole_sheet(tmp_path):
    import time
    import zipfile

    import openpyxl

    wb = openpyxl.Workbook()
    for row in (["品番", None, "単価"], ["型式"]):
        wb.active.append(row)
    wb.save(tmp_path / "a.xlsx")
    # 書式（行の高さ等）を最終行・最終列まで付けたブック。Excel は使用範囲を A1:XFD1048576 とし、書式だけの行を書く。
    rebuilt = tmp_path / "b.xlsx"
    with zipfile.ZipFile(tmp_path / "a.xlsx") as src, zipfile.ZipFile(rebuilt, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b'<dimension ref="A1:C2"', b'<dimension ref="A1:XFD1048576"')
                data = data.replace(b"</sheetData>", b'<row r="1048576" ht="20" customHeight="1"/></sheetData>')
            dst.writestr(item, data)
    start = time.monotonic()
    assert extract_segments(rebuilt)[0].text == "品番\t\t単価\n型式"
    # 使用範囲のとおりに埋めると数分かかる。
    assert time.monotonic() - start < 30


def test_extract_xlsx_formats_values_and_drops_blank_cells(tmp_path):
    import datetime

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([1, 1.5, 123456789012, 1e20, True, None, "  ", None])
    ws.append([datetime.date(2024, 1, 2), datetime.datetime(2024, 1, 2, 3, 4, 5), datetime.time(9, 30)])
    ws.append(["   ", None, " "])  # 空白だけの行は入れない
    ws["C5"] = "飛び"  # A1 以外から始まる値も列の位置を残す
    wb.save(tmp_path / "a.xlsx")
    assert extract_segments(tmp_path / "a.xlsx")[0].text.split("\n") == [
        "1\t1.5\t123456789012\t1e+20\tTrue",
        "2024-01-02\t2024-01-02 03:04:05\t09:30:00",
        "\t\t飛び",
    ]


def test_extract_xlsx_reads_far_apart_cells_without_calamine(tmp_path, monkeypatch):
    import zipfile

    import openpyxl

    from src.vector_db import extract

    wb = openpyxl.Workbook()
    wb.active["A1"] = "左上"
    wb.active["XFD2000"] = "右下"
    wb.save(tmp_path / "a.xlsx")
    # 値の入った範囲が上限を超えるブックは calamine に渡さない（使用範囲の無いブックは値のあるセルから数える）。
    called = []
    monkeypatch.setattr(extract, "_calamine_sheets", lambda source: called.append(source) or iter(()))
    no_dimension = tmp_path / "b.xlsx"
    with zipfile.ZipFile(tmp_path / "a.xlsx") as src, zipfile.ZipFile(no_dimension, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b'<dimension ref="A1:XFD2000"/>', b"")
            dst.writestr(item, data)
    # 使用範囲を実際より小さく（"A1" だけ）書くツールで作ったブックも、使用範囲を信じずに数える。
    understated = tmp_path / "c.xlsx"
    with zipfile.ZipFile(tmp_path / "a.xlsx") as src, zipfile.ZipFile(understated, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b'<dimension ref="A1:XFD2000"/>', b'<dimension ref="A1"/>')
            dst.writestr(item, data)
    for path in (tmp_path / "a.xlsx", no_dimension, understated):
        assert extract_segments(path)[0].text == "左上\n" + "\t" * 16383 + "右下"
    assert called == []


def test_extract_xlsx_falls_back_to_openpyxl_when_calamine_fails(tmp_path, monkeypatch):
    import openpyxl

    from src.vector_db import extract

    def broken(source):
        raise RuntimeError("calamine failed")
        yield

    wb = openpyxl.Workbook()
    wb.active.append(["品番", 10])
    wb.save(tmp_path / "a.xlsx")
    monkeypatch.setattr(extract, "_calamine_sheets", broken)
    assert extract_segments(tmp_path / "a.xlsx")[0].text == "品番\t10"


def test_extract_docx_reads_content_controls_and_tracked_insertions_in_order(tmp_path):
    import docx
    from docx.oxml import parse_xml

    w = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    d = docx.Document()
    d.add_paragraph("最初の段落")
    table = d.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "左"
    table.cell(0, 1).text = "右"
    # 本文レベルのコンテンツコントロール（社内テンプレートで多い）と、変更履歴で挿入した文字。
    d.element.body.insert(1, parse_xml(f"<w:sdt {w}><w:sdtContent><w:p><w:r><w:t>入力欄の文字</w:t></w:r></w:p></w:sdtContent></w:sdt>"))
    d.element.body.insert(
        len(d.element.body) - 1,  # 最後の要素（w:sectPr）の前
        parse_xml(f'<w:p {w}><w:r><w:t>前</w:t></w:r><w:ins w:id="1" w:author="a"><w:r><w:t>挿入</w:t></w:r></w:ins>'
                  f'<w:del w:id="2" w:author="a"><w:r><w:delText>削除</w:delText></w:r></w:del></w:p>'),
    )
    d.save(tmp_path / "a.docx")
    lines = [line for line in extract_segments(tmp_path / "a.docx")[0].text.split("\n") if line]
    assert lines == ["最初の段落", "入力欄の文字", "左\t右", "前挿入"]


def test_extract_pptx_reads_chart_titles(tmp_path):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["4月"]
    data.add_series("出荷", (1,))
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(4), Inches(3), data).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "月別の出荷数"
    prs.save(tmp_path / "a.pptx")
    assert extract_segments(tmp_path / "a.pptx")[0].text == "月別の出荷数"


def test_extract_rejects_unsupported_and_broken_files(tmp_path):
    (tmp_path / "a.bin").write_bytes(b"x")
    with pytest.raises(ExtractError):
        extract_segments(tmp_path / "a.bin")
    (tmp_path / "broken.docx").write_bytes(b"not a zip")
    with pytest.raises(ExtractError):
        extract_segments(tmp_path / "broken.docx")


def test_extract_reads_text_extensions_as_text(tmp_path):
    """text_extensions の拡張子は .txt と同じくテキストとして読み、中身がバイナリなら NotTextError にする。"""
    (tmp_path / "steel.mat").write_bytes("MATERIAL SPCC\nYOUNG 206000\n".encode("cp932"))
    with pytest.raises(ExtractError):
        extract_segments(tmp_path / "steel.mat")
    assert extract_segments(tmp_path / "steel.mat", text_extensions=(".mat",))[0].text == "MATERIAL SPCC\nYOUNG 206000\n"
    (tmp_path / "binary.mat").write_bytes(b"MATLAB 5.0\x00\x01\x02")
    with pytest.raises(NotTextError):
        extract_segments(tmp_path / "binary.mat", text_extensions=(".mat",))


# --- 埋め込み ---------------------------------------------------------------------


@pytest.fixture(autouse=True)
def _no_embedding_retry_wait(monkeypatch):
    """埋め込みの再試行の待ちを実際には眠らずに記録し、接続先の選び方の状態をテストごとに初期化する。"""
    sleeps: list[float] = []
    monkeypatch.setattr(embeddings, "_sleep", sleeps.append)
    monkeypatch.setattr(embeddings, "_round_robin_counters", {})
    monkeypatch.setattr(embeddings, "_failed_until", {})
    return sleeps


def _two_endpoints(*hosts: str) -> tuple[LLMEndpoint, ...]:
    return tuple(LLMEndpoint(base_url=f"http://{h}/v1", api_key="k", model="fake-embed") for h in hosts)


def _flaky_transport(server: FakeEmbeddingServer, failures: int, status: int = 503) -> tuple[httpx.MockTransport, list[str]]:
    """最初の failures 回は status を返し、その後は server に任せるトランスポート（送った先の URL を記録する）。"""
    calls: list[str] = []

    def handler(request: httpx.Request) -> httpx.Response:
        calls.append(str(request.url))
        if len(calls) <= failures:
            return httpx.Response(status, text="busy")
        return server.handler(request)

    return httpx.MockTransport(handler), calls


def test_embed_texts_batches_and_adds_prefix(tmp_path):
    server = FakeEmbeddingServer()
    cfg = make_config(tmp_path, vector_db_embedding_batch_size=2, vector_db_query_prefix="query: ", vector_db_document_prefix="passage: ")
    vectors = embed_texts(cfg, ["a", "b", "c"], kind="document", transport=server.transport)
    assert len(vectors) == 3 and len(vectors[0]) == DIM
    assert [r["input"] for r in server.requests] == [["passage: a", "passage: b"], ["passage: c"]]
    assert server.requests[0]["model"] == "fake-embed"
    embed_texts(cfg, ["q"], kind="query", transport=server.transport)
    assert server.requests[-1]["input"] == ["query: q"]


def test_embed_texts_fails_over_to_next_endpoint_on_connection_error(tmp_path, _no_embedding_retry_wait):
    server = FakeEmbeddingServer(fail_hosts=("down",))
    endpoints = _two_endpoints("down", "up")
    cfg = make_config(tmp_path, vector_db_embedding_endpoints=endpoints)
    embed_texts(cfg, ["a"], kind="query", transport=server.transport)
    assert server.requests[0]["url"] == "http://up/v1/embeddings"
    assert _no_embedding_retry_wait == []

    cfg = make_config(tmp_path, vector_db_embedding_endpoints=endpoints[:1], vector_db_embedding_max_retries=2)
    with pytest.raises(EmbeddingError, match="retried 2 times"):
        embed_texts(cfg, ["a"], kind="document", transport=server.transport)


def test_embed_texts_fails_over_to_next_endpoint_on_http_error(tmp_path, _no_embedding_retry_wait):
    """HTTP 200 以外が返ったら、同じリクエストを次の接続先へ送り直す（待たない）。"""
    server = FakeEmbeddingServer()
    transport, calls = _flaky_transport(server, failures=1, status=500)
    cfg = make_config(tmp_path, vector_db_embedding_endpoints=_two_endpoints("a", "b"))
    assert len(embed_texts(cfg, ["x"], kind="query", transport=transport)) == 1
    assert calls == ["http://a/v1/embeddings", "http://b/v1/embeddings"]
    assert _no_embedding_retry_wait == []


def test_embed_texts_retries_after_interval_until_success(tmp_path, _no_embedding_retry_wait):
    """全接続先でエラーなら embedding_retry_interval_seconds 待って全接続先を試し直す。"""
    server = FakeEmbeddingServer()
    transport, calls = _flaky_transport(server, failures=4)
    cfg = make_config(
        tmp_path,
        vector_db_embedding_endpoints=_two_endpoints("a", "b"),
        vector_db_embedding_max_retries=2,
        vector_db_embedding_retry_interval_seconds=1.0,
    )
    embed_texts(cfg, ["x"], kind="document", transport=transport)
    # 1巡目（a, b）・再試行1（a, b）が失敗し、再試行2の a で成功する
    assert len(calls) == 5
    assert sum(_no_embedding_retry_wait) == pytest.approx(2.0)


def test_embed_texts_gives_up_after_max_retries(tmp_path, _no_embedding_retry_wait):
    transport, calls = _flaky_transport(FakeEmbeddingServer(), failures=10**6)
    cfg = make_config(tmp_path, vector_db_embedding_max_retries=3, vector_db_embedding_retry_interval_seconds=2.0)
    with pytest.raises(EmbeddingError, match="retried 3 times.*HTTP 503"):
        embed_texts(cfg, ["x"], kind="document", transport=transport)
    assert len(calls) == 4
    assert sum(_no_embedding_retry_wait) == pytest.approx(6.0)


def test_embed_texts_retries_without_limit_when_max_retries_is_zero(tmp_path, _no_embedding_retry_wait):
    transport, calls = _flaky_transport(FakeEmbeddingServer(), failures=25)
    cfg = make_config(tmp_path, vector_db_embedding_max_retries=0, vector_db_embedding_retry_interval_seconds=0.5)
    embed_texts(cfg, ["x"], kind="document", transport=transport)
    assert len(calls) == 26
    assert sum(_no_embedding_retry_wait) == pytest.approx(12.5)


def test_search_embedding_uses_its_own_bounded_retries(tmp_path, _no_embedding_retry_wait):
    """検索語の埋め込みは embedding_max_retries = 0（無制限）でも search_embedding_max_retries 回でやめる。"""
    cfg = make_config(tmp_path, vector_db_embedding_max_retries=0, vector_db_search_embedding_max_retries=2)
    transport, calls = _flaky_transport(FakeEmbeddingServer(), failures=10**6)
    with pytest.raises(EmbeddingError, match="retried 2 times"):
        embed_texts(cfg, ["x"], kind="query", transport=transport)
    assert len(calls) == 3
    assert sum(_no_embedding_retry_wait) == pytest.approx(20.0)

    # 0 は「試し直さない」
    _no_embedding_retry_wait.clear()
    cfg = make_config(tmp_path, vector_db_search_embedding_max_retries=0)
    transport, calls = _flaky_transport(FakeEmbeddingServer(), failures=10**6)
    with pytest.raises(EmbeddingError):
        embed_texts(cfg, ["x"], kind="query", transport=transport)
    assert len(calls) == 1 and _no_embedding_retry_wait == []


def test_embed_texts_does_not_retry_client_errors(tmp_path, _no_embedding_retry_wait):
    """408・429 以外の 4xx は待っても直らないため、全接続先を1巡したら試し直さない。"""
    transport, calls = _flaky_transport(FakeEmbeddingServer(), failures=10**6, status=400)
    cfg = make_config(tmp_path, vector_db_embedding_endpoints=_two_endpoints("a", "b"), vector_db_embedding_max_retries=0)
    with pytest.raises(EmbeddingError, match="HTTP 400"):
        embed_texts(cfg, ["x"], kind="document", transport=transport)
    assert len(calls) == 2 and _no_embedding_retry_wait == []

    transport, calls = _flaky_transport(FakeEmbeddingServer(), failures=2, status=429)
    embed_texts(cfg, ["x"], kind="document", transport=transport)
    assert len(calls) == 3


def test_embed_retry_wait_stops_when_requested(tmp_path):
    transport, calls = _flaky_transport(FakeEmbeddingServer(), failures=10**6)
    cfg = make_config(tmp_path, vector_db_embedding_max_retries=0)
    checks = iter([False, False, True])
    with pytest.raises(EmbeddingStopped):
        embed_images(cfg, [b"png"], transport=transport, should_stop=lambda: next(checks))
    assert len(calls) == 1


@pytest.mark.parametrize(
    ("strategy", "expected_first"),
    [("priority_failover", ["a", "a", "a"]), ("round_robin", ["a", "b", "c"])],
)
def test_embed_texts_routing_strategy(tmp_path, strategy, expected_first):
    server = FakeEmbeddingServer()
    cfg = make_config(tmp_path, vector_db_embedding_endpoints=_two_endpoints("a", "b", "c"), vector_db_embedding_routing_strategy=strategy)
    for _ in range(3):
        embed_texts(cfg, ["x"], kind="query", transport=server.transport)
    assert [httpx.URL(r["url"]).host for r in server.requests] == expected_first


def test_embed_texts_random_strategy_uses_all_endpoints(tmp_path, monkeypatch):
    server = FakeEmbeddingServer()
    monkeypatch.setattr(embeddings.random, "sample", lambda seq, k: list(reversed(seq)))
    cfg = make_config(tmp_path, vector_db_embedding_endpoints=_two_endpoints("a", "b"), vector_db_embedding_routing_strategy="random")
    embed_texts(cfg, ["x"], kind="query", transport=server.transport)
    assert server.requests[0]["url"] == "http://b/v1/embeddings"


def test_priority_failover_puts_failed_endpoint_last_for_a_while(tmp_path):
    """priority_failover では、エラーが返った接続先を一定時間後回しにする（毎回タイムアウトを待たない）。"""
    server = FakeEmbeddingServer()
    transport, calls = _flaky_transport(server, failures=1)
    cfg = make_config(tmp_path, vector_db_embedding_endpoints=_two_endpoints("a", "b"))
    embed_texts(cfg, ["x"], kind="query", transport=transport)
    embed_texts(cfg, ["y"], kind="query", transport=transport)
    assert calls == ["http://a/v1/embeddings", "http://b/v1/embeddings", "http://b/v1/embeddings"]


def test_embed_texts_rejects_token_level_embeddings(tmp_path):
    server = FakeEmbeddingServer(two_dimensional=True)
    with pytest.raises(EmbeddingError, match="pooling"):
        embed_texts(make_config(tmp_path), ["a"], kind="query", transport=server.transport)


def test_embed_texts_reports_http_error(tmp_path):
    transport = httpx.MockTransport(lambda r: httpx.Response(500, text="model not loaded"))
    with pytest.raises(EmbeddingError, match="HTTP 500.*model not loaded"):
        embed_texts(make_config(tmp_path), ["a"], kind="query", transport=transport)


# --- 保存先 ----------------------------------------------------------------------


def _exercise_store(store):
    store.ensure_collection("docs", DIM)
    store.ensure_collection("docs", DIM)  # 2回目は何もしない
    store.replace_source("docs", "/a.md", [ChunkRecord("p.1", "溶接の基準"), ChunkRecord("p.2", "検査の方法")], [fake_vector("溶接の基準"), fake_vector("検査の方法")])
    store.replace_source("docs", "/b.md", [ChunkRecord("", "会議室の予約")], [fake_vector("会議室の予約")])
    hits = store.search("docs", fake_vector("溶接の基準"), 2)
    assert hits[0].source_path == "/a.md" and hits[0].locator == "p.1" and hits[0].text == "溶接の基準"
    assert hits[0].score == pytest.approx(1.0, abs=1e-5)
    # 入れ替えると古いチャンクは残らない
    store.replace_source("docs", "/a.md", [ChunkRecord("", "差し替え後")], [fake_vector("差し替え後")])
    texts = {h.text for h in store.search("docs", fake_vector("x"), 10)}
    assert texts == {"差し替え後", "会議室の予約"}
    store.delete_source("docs", "/b.md")
    assert {h.source_path for h in store.search("docs", fake_vector("x"), 10)} == {"/a.md"}
    with pytest.raises(StoreError):
        store.ensure_collection("docs", DIM + 1)
    assert store.search("missing", fake_vector("x"), 3) == []
    store.drop_collection("docs")
    store.drop_collection("docs")
    assert store.search("docs", fake_vector("x"), 3) == []


@pytest.mark.parametrize("distance", ["Cosine", "Euclid", "Manhattan"])
def test_sqlite_vec_store(tmp_path, distance):
    store = SqliteVecStore(tmp_path, distance=distance)
    try:
        # 完全一致の score が 1.0 になる（Euclid・Manhattan は 1 / (1 + 距離) に換算）
        _exercise_store(store)
    finally:
        store.close()


def test_sqlite_vec_store_rejects_dot(tmp_path):
    with pytest.raises(StoreError, match="Dot"):
        SqliteVecStore(tmp_path, distance="Dot")


def test_sqlite_vec_store_rejects_distance_mismatch(tmp_path):
    store = SqliteVecStore(tmp_path, distance="Euclid")
    try:
        store.ensure_collection("docs", DIM)
        store.replace_source("docs", "/a.md", [ChunkRecord("", "溶接の基準"), ChunkRecord("", "会議室の予約")], [fake_vector("溶接の基準"), fake_vector("会議室の予約")])
    finally:
        store.close()
    other = SqliteVecStore(tmp_path, distance="Cosine")
    try:
        with pytest.raises(StoreError, match="distance.*Rebuild"):
            other.ensure_collection("docs", DIM)
        # 検索は作ったときの距離関数（Euclid）で換算する
        hits = other.search("docs", fake_vector("溶接の基準"), 2)
        assert [h.text for h in hits] == ["溶接の基準", "会議室の予約"]
        assert 0.0 < hits[1].score < hits[0].score == pytest.approx(1.0)
        # 再構築（消して作り直す）すれば新しい距離関数で作れる
        other.drop_collection("docs")
        other.ensure_collection("docs", DIM)
    finally:
        other.close()


def test_sqlite_vec_store_migrates_db_without_distance_column(tmp_path):
    import sqlite3

    import sqlite_vec

    conn = sqlite3.connect(str(tmp_path / "vectors.sqlite"))
    conn.enable_load_extension(True)
    sqlite_vec.load(conn)
    conn.executescript(
        """
        CREATE TABLE vec_collections (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL UNIQUE, dim INTEGER NOT NULL);
        CREATE TABLE chunks (id INTEGER PRIMARY KEY AUTOINCREMENT, collection_id INTEGER NOT NULL, source_path TEXT NOT NULL,
                             locator TEXT NOT NULL, text TEXT NOT NULL);
        """
    )
    conn.execute("INSERT INTO vec_collections (name, dim) VALUES ('docs', ?)", (DIM,))
    conn.execute(f"CREATE VIRTUAL TABLE vec_c1 USING vec0(embedding float[{DIM}] distance_metric=cosine)")
    conn.commit()
    conn.close()
    # 旧版の DB はコサイン距離で作られているため、そのまま Cosine として使える
    store = SqliteVecStore(tmp_path)
    try:
        store.ensure_collection("docs", DIM)
        store.replace_source("docs", "/a.md", [ChunkRecord("", "溶接の基準")], [fake_vector("溶接の基準")])
        assert store.search("docs", fake_vector("溶接の基準"), 1)[0].score == pytest.approx(1.0, abs=1e-5)
    finally:
        store.close()


@pytest.mark.parametrize("distance", ["Cosine", "Euclid", "Dot", "Manhattan"])
def test_qdrant_store_uses_prefixed_collection_names(distance):
    fake = FakeQdrant()
    store = QdrantStore("http://qdrant:6333/", prefix="locohane_x_", api_key="secret", distance=distance, transport=fake.transport)
    try:
        # 完全一致の score が 1.0 になる（Euclid・Manhattan は 1 / (1 + 距離) に換算）
        _exercise_store(store)
    finally:
        store.close()
    assert all(path.startswith("/collections/locohane_x_") for _, path in fake.calls)


def test_qdrant_store_creates_collection_with_distance_and_rejects_mismatch():
    fake = FakeQdrant()
    store = QdrantStore("http://qdrant:6333", prefix="p_", distance="Euclid", transport=fake.transport)
    try:
        store.ensure_collection("docs", DIM)
        assert fake.collections["p_docs"]["distance"] == "Euclid"
        store.replace_source("docs", "/a.md", [ChunkRecord("", "溶接の基準"), ChunkRecord("", "会議室の予約")], [fake_vector("溶接の基準"), fake_vector("会議室の予約")])
        hits = store.search("docs", fake_vector("溶接の基準"), 2)
        # 距離が小さい順（換算後の score が大きい順）
        assert [h.text for h in hits] == ["溶接の基準", "会議室の予約"]
        assert 0.0 < hits[1].score < hits[0].score == pytest.approx(1.0)
    finally:
        store.close()
    # 距離関数を変えた設定では、既存のコレクションへの取り込みを拒否し再構築を求める
    other = QdrantStore("http://qdrant:6333", prefix="p_", distance="Cosine", transport=fake.transport)
    try:
        with pytest.raises(StoreError, match="distance.*Rebuild"):
            other.ensure_collection("docs", DIM)
        # 検索は作ったときの距離関数（Euclid）で換算する
        assert other.search("docs", fake_vector("溶接の基準"), 1)[0].score == pytest.approx(1.0)
    finally:
        other.close()


def test_qdrant_store_without_preset_keeps_qdrant_defaults():
    fake = FakeQdrant()
    store = QdrantStore("http://qdrant:6333", prefix="p_", transport=fake.transport)
    try:
        store.ensure_collection("docs", DIM)
        store.apply_preset("docs")
        store.search("docs", fake_vector("溶接"), 1)
    finally:
        store.close()
    config = fake.collections["p_docs"]["config"]
    assert config["quantization_config"] is None and "on_disk" not in config["params"]["vectors"]
    assert fake.patches == [] and "params" not in fake.queries[0]


@pytest.mark.parametrize("preset", ["1m", "2m", "3m", "4m", "5m", "6m", "7m", "8m", "9m", "10m"])
def test_qdrant_store_creates_collection_with_preset(preset):
    fake = FakeQdrant()
    store = QdrantStore("http://qdrant:6333", prefix="p_", preset=preset, transport=fake.transport)
    expected = QDRANT_PRESETS[preset]
    try:
        _exercise_store(store)
        # 作ったときに設定したので、もう一度 ensure・apply しても PATCH しない
        store.ensure_collection("docs", DIM)
        store.apply_preset("docs")
    finally:
        store.close()
    config = fake.collections["p_docs"]["config"]
    assert config["hnsw_config"]["m"] == expected.m and config["hnsw_config"]["ef_construct"] == expected.ef_construct
    assert config["params"]["vectors"]["on_disk"] is expected.vectors_on_disk and config["params"]["on_disk_payload"] is True
    assert config["optimizer_config"]["default_segment_number"] == expected.segments
    assert config["quantization_config"]["scalar"]["type"] == "int8"
    assert fake.patches == []
    assert fake.queries[-1]["params"] == {"hnsw_ef": expected.hnsw_ef, "quantization": {"rescore": True, "oversampling": expected.oversampling}}


def test_qdrant_presets_grow_with_scale_and_match_config_choices():
    from src.config import VECTOR_DB_QDRANT_PRESETS

    assert tuple(QDRANT_PRESETS) == VECTOR_DB_QDRANT_PRESETS[1:] and VECTOR_DB_QDRANT_PRESETS[0] == "none"
    presets = list(QDRANT_PRESETS.values())
    for smaller, larger in zip(presets, presets[1:]):
        assert larger.m >= smaller.m and larger.ef_construct >= smaller.ef_construct and larger.hnsw_ef > smaller.hnsw_ef


def test_qdrant_store_applies_preset_to_existing_collection():
    fake = FakeQdrant()
    old = QdrantStore("http://qdrant:6333", prefix="p_", transport=fake.transport)
    try:
        old.ensure_collection("docs", DIM)
        old.replace_source("docs", "/a.md", [ChunkRecord("", "溶接の基準")], [fake_vector("溶接の基準")])
    finally:
        old.close()
    store = QdrantStore("http://qdrant:6333", prefix="p_", preset="3m", transport=fake.transport)
    try:
        # 無いコレクションは何もしない
        store.apply_preset("missing")
        store.apply_preset("docs")
        store.apply_preset("docs")
        assert store.search("docs", fake_vector("溶接の基準"), 1)[0].text == "溶接の基準"
    finally:
        store.close()
    # 設定が違うときだけ1回 PATCH し、ベクターは消さない（埋め込み直し不要）
    assert [name for name, _ in fake.patches] == ["p_docs"]
    body = fake.patches[0][1]
    assert body["vectors"] == {"": {"on_disk": True}} and body["hnsw_config"]["m"] == QDRANT_PRESETS["3m"].m
    assert body["params"] == {"on_disk_payload": True}
    assert len(fake.collections["p_docs"]["points"]) == 1


def test_ingest_applies_qdrant_preset_without_changed_files(tmp_path):
    docs = tmp_path / "docs"
    docs.mkdir()
    (docs / "a.md").write_text("溶接の基準は JIS に従う。", encoding="utf-8")
    embed, qdrant = FakeEmbeddingServer(), FakeQdrant()
    cfg = make_config(tmp_path, vector_db_backend="qdrant")
    with vcat.open_catalog(cfg.vector_db_dir) as catalog:
        catalog.create("docs", str(docs))
    ingest.ingest_collection(cfg, "docs", embedding_transport=embed.transport, store_transport=qdrant.transport)
    assert qdrant.patches == []
    tuned = make_config(tmp_path, vector_db_backend="qdrant", vector_db_qdrant_preset="2m")
    result = ingest.ingest_collection(tuned, "docs", embedding_transport=embed.transport, store_transport=qdrant.transport)
    assert result.unchanged == 1
    assert [name for name, _ in qdrant.patches] == ["locohane_test_docs", "locohane_test_docs.path"]


def test_qdrant_store_sends_api_key_and_reports_connection_error():
    seen = {}

    def handler(request):
        seen["key"] = request.headers.get("api-key")
        raise httpx.ConnectError("refused", request=request)

    store = QdrantStore("http://qdrant:6333", prefix="p_", api_key="secret", transport=httpx.MockTransport(handler))
    with pytest.raises(StoreError, match="Cannot connect"):
        store.search("docs", [0.1], 1)
    assert seen["key"] == "secret"


# --- カタログ --------------------------------------------------------------------


def test_catalog_validates_names_and_duplicates(tmp_path):
    with vcat.open_catalog(tmp_path) as catalog:
        with pytest.raises(vcat.CatalogError):
            catalog.create("日本語", "C:/x")
        catalog.create("docs", "C:/x", "説明")
        with pytest.raises(vcat.CatalogError):
            catalog.create("docs", "C:/y")
        updated = catalog.update("docs", description="新しい説明", recursive=False)
        assert (updated.description, updated.recursive) == ("新しい説明", False)
        catalog.delete("docs")
        assert catalog.list_collections() == []


# --- 取り込みと検索 ---------------------------------------------------------------


@pytest.fixture(params=["sqlite_vec", "qdrant"])
def env(request, tmp_path):
    docs = tmp_path / "docs"
    (docs / "sub").mkdir(parents=True)
    (docs / "weld.md").write_text("溶接ビードの幅は5mm以上とする。", encoding="utf-8")
    (docs / "sub" / "room.txt").write_text("会議室の予約はポータルから行う。", encoding="utf-8")
    (docs / "~$lock.docx").write_bytes(b"office lock file")
    (docs / "image.xyz").write_bytes(b"unsupported")
    cfg = make_config(tmp_path, vector_db_backend=request.param)
    embed = FakeEmbeddingServer()
    qdrant = FakeQdrant()
    kwargs = {"embedding_transport": embed.transport, "store_transport": qdrant.transport}
    with vcat.open_catalog(cfg.vector_db_dir) as catalog:
        catalog.create("docs", str(docs), "テスト文書")
    return {"cfg": cfg, "docs": docs, "kwargs": kwargs, "embed": embed, "qdrant": qdrant}


def test_ingest_is_incremental(env):
    cfg, docs, kwargs = env["cfg"], env["docs"], env["kwargs"]
    progress = []
    first = ingest.ingest_collection(cfg, "docs", progress=lambda *a: progress.append(a), **kwargs)
    assert (first.added, first.unchanged, first.failed) == (2, 0, [])
    assert progress[-1] == (2, 2, "")
    # 記録が終わるたびにも件数を通知する（最後のファイルを読み始めた後も、終了の通知を待たずに 2 / 2 になる）。
    assert progress[-2][:2] == (2, 2) and progress[-2][2]
    requests_after_first = len(env["embed"].requests)

    second = ingest.ingest_collection(cfg, "docs", **kwargs)
    assert (second.added, second.updated, second.unchanged) == (0, 0, 2)
    assert len(env["embed"].requests) == requests_after_first  # 変更が無ければ埋め込み直さない

    # 更新日時だけ変わって内容が同じなら埋め込み直さない
    weld = docs / "weld.md"
    os.utime(weld, (weld.stat().st_atime, weld.stat().st_mtime + 10))
    assert ingest.ingest_collection(cfg, "docs", **kwargs).unchanged == 2
    assert len(env["embed"].requests) == requests_after_first

    weld.write_text("溶接ビードの幅は6mm以上に改訂した。", encoding="utf-8")
    (docs / "sub" / "room.txt").unlink()
    (docs / "empty.md").write_text("   ", encoding="utf-8")
    third = ingest.ingest_collection(cfg, "docs", **kwargs)
    assert (third.updated, third.removed) == (1, 1)
    # 本文を取り出せないファイルも、失敗にせずファイルパスだけを登録する（ファイル名で探せるように）。
    assert (third.added, third.path_only, third.failed) == (1, 1, [])

    result = search.search(cfg, "溶接ビードの幅", **kwargs)
    assert [h.text for h in result.hits] == ["溶接ビードの幅は6mm以上に改訂した。"]
    [collection] = search.list_collections(cfg)
    assert (collection.file_count, collection.chunk_count, collection.model, collection.dim) == (2, 1, "fake-embed", DIM)
    assert str(docs / "empty.md") in {h.source_path for h in search.search_paths(cfg, "empty", **kwargs).hits}


def test_ingest_reuses_one_embedding_client(env, monkeypatch):
    # ファイルごとに接続し直すと、Windows の localhost では毎回約2秒かかるため、取り込み全体で1つを使い回す。
    cfg, kwargs = env["cfg"], env["kwargs"]
    opened, closed = [], []
    real_open = ingest.open_client

    def counting_open(*args, **kw):
        client = real_open(*args, **kw)
        opened.append(client)
        real_close = client.close
        client.close = lambda: (closed.append(client), real_close())[1]
        return client

    monkeypatch.setattr(ingest, "open_client", counting_open)
    result = ingest.ingest_collection(cfg, "docs", **kwargs)
    assert result.added == 2 and len(env["embed"].requests) == 2
    assert len(opened) == 1 and closed == opened


def test_embed_texts_does_not_close_given_client(tmp_path):
    server = FakeEmbeddingServer()
    cfg = make_config(tmp_path)
    with open_client(cfg, transport=server.transport) as client:
        embed_texts(cfg, ["a"], kind="document", client=client)
        embed_texts(cfg, ["b"], kind="document", client=client)
        assert not client.is_closed
    assert len(server.requests) == 2


def test_shared_client_is_reused_until_closed(tmp_path):
    cfg = make_config(tmp_path)
    client = shared_client(cfg)
    assert shared_client(cfg) is client
    client.close()
    reopened = shared_client(cfg)
    assert reopened is not client and not reopened.is_closed


def test_qdrant_shared_store_keeps_client_open(monkeypatch):
    from src.vector_db import qdrant_store

    monkeypatch.setattr(qdrant_store, "_shared_clients", {})
    first = QdrantStore("http://qdrant:6333/", prefix="p_", shared=True)
    first.close()
    second = QdrantStore("http://qdrant:6333", prefix="p_", shared=True)
    assert second._client is first._client and not first._client.is_closed
    own = QdrantStore("http://qdrant:6333", prefix="p_")
    own.close()
    assert own._client.is_closed and not first._client.is_closed


def test_ingest_skips_files_over_size_limit_and_non_recursive(env):
    cfg = make_config(env["cfg"].vector_db_dir.parent, vector_db_backend=env["cfg"].vector_db_backend, vector_db_max_file_mb=0.00001)
    result = ingest.ingest_collection(cfg, "docs", **env["kwargs"])
    assert len(result.failed) == 2 and "limit" in result.failed[0][1]
    assert result.too_large == 2 and result.to_dict()["too_large"] == 2
    with vcat.open_catalog(cfg.vector_db_dir) as catalog:
        catalog.update("docs", recursive=False)
    files = ingest.list_source_files(env["docs"], False, cfg.vector_db_include_extensions)
    assert [p.name for p in files] == ["weld.md"]


def test_ingest_reads_text_extensions_only_when_included(env):
    """text_extensions の拡張子は include_extensions にも書いたときだけ取り込み、バイナリならパスだけを登録する。"""
    docs, kwargs = env["docs"], env["kwargs"]
    (docs / "spcc.mat").write_text("MATERIAL SPCC 引張強さ 270MPa", encoding="utf-8")
    (docs / "binary.mat").write_bytes(b"\x00\x01binary")
    (docs / "part.step").write_text("ISO-10303-21;", encoding="ascii")
    base = env["cfg"]
    # text_extensions に書いても include_extensions に無ければ取り込まない（従来どおり）。
    only_text = make_config(base.vector_db_dir.parent, vector_db_backend=base.vector_db_backend, vector_db_text_extensions=(".mat",))
    assert ingest.ingest_collection(only_text, "docs", **kwargs).added == 2

    cfg = make_config(
        base.vector_db_dir.parent, vector_db_backend=base.vector_db_backend,
        vector_db_include_extensions=(*base.vector_db_include_extensions, ".mat", ".step"), vector_db_text_extensions=(".mat",),
    )
    result = ingest.ingest_collection(cfg, "docs", **kwargs)
    # .step は text_extensions に無いため対象外のまま。バイナリの .mat は失敗にせずパスだけを登録する。
    assert (result.added, result.path_only, result.failed) == (2, 1, [])
    hits = search.search(cfg, "SPCC 引張強さ", **kwargs).hits
    assert any(h.text == "MATERIAL SPCC 引張強さ 270MPa" for h in hits)


def test_ingest_requires_rebuild_when_embedding_model_changes(env):
    cfg, kwargs = env["cfg"], env["kwargs"]
    ingest.ingest_collection(cfg, "docs", **kwargs)
    other = make_config(
        cfg.vector_db_dir.parent,
        vector_db_backend=cfg.vector_db_backend,
        vector_db_embedding_endpoints=(LLMEndpoint(base_url="http://embed/v1", api_key="k", model="other-embed"),),
    )
    with pytest.raises(ingest.IngestError, match="rebuild"):
        ingest.ingest_collection(other, "docs", **kwargs)
    # 検索も対象外になり、理由が分かる
    with pytest.raises(search.SearchError, match="other-embed"):
        search.search(other, "溶接", **kwargs)
    rebuilt = ingest.ingest_collection(other, "docs", rebuild=True, **kwargs)
    assert rebuilt.added == 2
    assert search.list_collections(other)[0].model == "other-embed"


def test_ingest_requires_rebuild_when_document_prefix_changes(env):
    cfg, kwargs = env["cfg"], env["kwargs"]
    ingest.ingest_collection(cfg, "docs", **kwargs)
    assert search.list_collections(cfg)[0].document_prefix == cfg.vector_db_document_prefix
    other = make_config(cfg.vector_db_dir.parent, vector_db_backend=cfg.vector_db_backend, vector_db_document_prefix="passage: ")
    with pytest.raises(ingest.IngestError, match="document_prefix"):
        ingest.ingest_collection(other, "docs", **kwargs)
    with pytest.raises(search.SearchError, match="document_prefix"):
        search.search(other, "溶接", **kwargs)
    # query_prefix だけの違いは再構築不要
    query_only = make_config(cfg.vector_db_dir.parent, vector_db_backend=cfg.vector_db_backend, vector_db_query_prefix="query: ")
    assert search.search(query_only, "溶接", **kwargs).hits
    ingest.ingest_collection(other, "docs", rebuild=True, **kwargs)
    assert search.list_collections(other)[0].document_prefix == "passage: "
    assert search.search(other, "溶接", **kwargs).hits


def test_catalog_adds_document_prefix_column_to_old_db(tmp_path):
    import sqlite3

    conn = sqlite3.connect(str(tmp_path / vcat.CATALOG_FILENAME))
    conn.executescript(
        """
        CREATE TABLE collections (name TEXT PRIMARY KEY, description TEXT NOT NULL DEFAULT '', source_folder TEXT NOT NULL,
            recursive INTEGER NOT NULL DEFAULT 1, model TEXT, dim INTEGER, created_at TEXT NOT NULL, last_ingested_at TEXT);
        INSERT INTO collections (name, source_folder, model, dim, created_at) VALUES ('old', 'C:/docs', 'fake-embed', 8, '2026-01-01');
        """
    )
    conn.close()
    with vcat.open_catalog(tmp_path) as catalog:
        assert catalog.get("old").document_prefix is None


def test_ingest_stops_without_deleting_unprocessed_files(env):
    cfg, kwargs = env["cfg"], env["kwargs"]
    ingest.ingest_collection(cfg, "docs", **kwargs)
    result = ingest.ingest_collection(cfg, "docs", should_stop=lambda: True, **kwargs)
    assert result.removed == 0
    assert search.list_collections(cfg)[0].file_count == 2


def test_ingest_runs_files_in_parallel_with_ingest_workers(env):
    """ingest_workers > 1 なら複数ファイルの埋め込みを同時に行い、結果は1つずつ処理したときと同じになる。"""
    import threading
    import time

    cfg, docs = env["cfg"], env["docs"]
    for i in range(6):
        (docs / f"extra{i}.md").write_text(f"追加の文書{i}。手順{i}を守る。", encoding="utf-8")
    (docs / "empty.md").write_text("   ", encoding="utf-8")
    embed = env["embed"]
    lock = threading.Lock()
    state = {"active": 0, "peak": 0}

    def slow_handler(request):
        with lock:
            state["active"] += 1
            state["peak"] = max(state["peak"], state["active"])
        try:
            time.sleep(0.05)
            return embed.handler(request)
        finally:
            with lock:
                state["active"] -= 1

    kwargs = {**env["kwargs"], "embedding_transport": httpx.MockTransport(slow_handler)}
    parallel = make_config(cfg.vector_db_dir.parent, vector_db_backend=cfg.vector_db_backend, vector_db_ingest_workers=3)
    progress = []
    result = ingest.ingest_collection(parallel, "docs", progress=lambda *a: progress.append(a), **kwargs)
    assert (result.added, result.unchanged, result.path_only, result.failed) == (9, 0, 1, [])
    assert 1 < state["peak"] <= 3
    assert progress[-1] == (9, 9, "")
    assert search.list_collections(parallel)[0].file_count == 9
    assert [h.text for h in search.search(parallel, "溶接ビードの幅", **kwargs).hits][0] == "溶接ビードの幅は5mm以上とする。"
    # 2回目は差分なし（並列でも記録が正しく残っている）
    again = ingest.ingest_collection(parallel, "docs", **kwargs)
    assert (again.added, again.updated, again.unchanged) == (0, 0, 9)


def test_ingest_in_parallel_stops_on_embedding_error(env):
    cfg = make_config(env["cfg"].vector_db_dir.parent, vector_db_backend=env["cfg"].vector_db_backend, vector_db_ingest_workers=2)
    failing = FakeEmbeddingServer(fail_hosts=("embed",))
    kwargs = {**env["kwargs"], "embedding_transport": failing.transport}
    with pytest.raises(EmbeddingError):
        ingest.ingest_collection(cfg, "docs", **kwargs)
    assert search.list_collections(cfg)[0].file_count == 0


def test_ingest_stop_during_embedding_retry_keeps_file_for_next_run(env):
    """埋め込みの再試行を待っている間に停止したら、そのファイルを記録せずに止め、次回の取り込みで処理する。"""
    cfg = make_config(env["cfg"].vector_db_dir.parent, vector_db_backend=env["cfg"].vector_db_backend, vector_db_embedding_max_retries=0)
    stop = {"requested": False}

    def down(request: httpx.Request) -> httpx.Response:
        stop["requested"] = True
        return httpx.Response(503, text="loading")

    kwargs = {**env["kwargs"], "embedding_transport": httpx.MockTransport(down)}
    result = ingest.ingest_collection(cfg, "docs", should_stop=lambda: stop["requested"], **kwargs)
    assert result.added == 0
    assert search.list_collections(cfg)[0].file_count == 0

    again = ingest.ingest_collection(cfg, "docs", **env["kwargs"])
    assert again.added > 0 and again.removed == 0


def test_ingest_stop_reports_processed_count(env):
    """埋め込みの途中で停止しても、最後の進捗は記録まで終わったファイルの件数を指す。

    以前は停止すると最後の通知を送らず、最後に読み始めたファイルの時点の件数（ここでは0件）のまま残っていた。
    """
    cfg, kwargs = env["cfg"], env["kwargs"]
    progress = []
    result = ingest.ingest_collection(
        cfg, "docs", progress=lambda *a: progress.append(a), should_stop=lambda: len(env["embed"].requests) >= 1, **kwargs
    )
    assert 1 <= result.added < 2
    assert progress[-1] == (result.added, 2, "")


def test_ingest_errors_for_missing_collection_or_folder(env):
    with pytest.raises(ingest.IngestError):
        ingest.ingest_collection(env["cfg"], "missing", **env["kwargs"])
    with vcat.open_catalog(env["cfg"].vector_db_dir) as catalog:
        catalog.update("docs", source_folders=[str(env["docs"]), str(env["docs"] / "nothing")])
    # 1つでも見えないフォルダがあれば、そのフォルダのファイルを消さないよう取り込まない。
    with pytest.raises(ingest.IngestError, match="folder"):
        ingest.ingest_collection(env["cfg"], "docs", **env["kwargs"])


def test_ingest_multiple_source_folders(env, tmp_path):
    cfg, kwargs = env["cfg"], env["kwargs"]
    extra = tmp_path / "extra"
    (extra / "sub").mkdir(parents=True)
    (extra / "sub" / "guide.md").write_text("出張の申請は3日前までに行う。", encoding="utf-8")
    with vcat.open_catalog(cfg.vector_db_dir) as catalog:
        # 文字列なら改行区切り。空行と同じフォルダの重複は除く。
        catalog.update("docs", source_folders=f"{env['docs']}\n\n{extra}\n{env['docs']}")
        assert catalog.get("docs").source_folders == (str(env["docs"]), str(extra))
    assert ingest.ingest_collection(cfg, "docs", **kwargs).added == 3
    with vcat.open_catalog(cfg.vector_db_dir) as catalog:
        files = catalog.files("docs")
    assert str(extra / "sub" / "guide.md") in files
    hits = search.search(cfg, "出張の申請", top_k=1, **kwargs).hits
    assert hits[0].source_path == str(extra / "sub" / "guide.md")

    # フォルダを1つ外すと、そのフォルダのファイルは次の取り込みで消える。
    with vcat.open_catalog(cfg.vector_db_dir) as catalog:
        catalog.update("docs", source_folders=[str(env["docs"])])
    assert ingest.ingest_collection(cfg, "docs", **kwargs).removed == 1


def test_scan_skips_source_folder_inside_another_when_recursive(tmp_path):
    (tmp_path / "sub").mkdir()
    (tmp_path / "a.md").write_text("a", encoding="utf-8")
    (tmp_path / "sub" / "b.md").write_text("b", encoding="utf-8")
    other = tmp_path.parent / f"{tmp_path.name}_other"
    other.mkdir()
    (other / "c.md").write_text("c", encoding="utf-8")
    for workers in (1, 4):
        files = [p for p, _ in ingest.scan_source_files([tmp_path / "sub", tmp_path, other], True, (".md",), workers=workers).files]
        assert files == sorted([tmp_path / "a.md", tmp_path / "sub" / "b.md", other / "c.md"], key=lambda p: str(p).lower())
    # 直下だけなら入れ子のフォルダもそれぞれの直下を取り込む。
    files = ingest.list_source_files([tmp_path, tmp_path / "sub"], False, (".md",))
    assert sorted(files) == sorted([tmp_path / "a.md", tmp_path / "sub" / "b.md"])
    assert ingest.relative_path_text(ingest.source_root([tmp_path, tmp_path / "sub"], tmp_path / "sub" / "b.md"), tmp_path / "sub" / "b.md") == "b.md"


def test_scan_skips_paths_matching_exclude_keywords(tmp_path):
    for rel in ("a.md", "Old/b.md", "sub/old_c.md", "sub/d.md", "sub/older/e.md", "x/old/f.md"):
        (tmp_path / rel).parent.mkdir(parents=True, exist_ok=True)
        (tmp_path / rel).write_text("x", encoding="utf-8")
    for workers in (1, 4):
        # 大文字小文字は区別せず、フォルダ名・ファイル名・途中のパスのどこに含んでも除外する。
        files = [p for p, _ in ingest.scan_source_files(tmp_path, True, (".md",), workers=workers, exclude_keywords=["OLD"]).files]
        assert files == [tmp_path / "a.md", tmp_path / "sub" / "d.md"]
    # / で囲むと名前がちょうどそれのフォルダだけ。区切りは \ でもよい。
    files = ingest.list_source_files(tmp_path, True, (".md",), exclude_keywords=["/old/", "sub\\d"])
    assert sorted(files) == sorted([tmp_path / "a.md", tmp_path / "sub" / "old_c.md", tmp_path / "sub" / "older" / "e.md"])
    # "*." で始まれば拡張子の指定（ファイル名の末尾だけ。名前の途中・フォルダには当てない）。
    for rel in ("g.MD", "h.md.txt", "y.md/i.txt"):
        (tmp_path / rel).parent.mkdir(parents=True, exist_ok=True)
        (tmp_path / rel).write_text("x", encoding="utf-8")
    files = ingest.list_source_files(tmp_path, True, (".md", ".txt"), exclude_keywords=["*.md"])
    assert sorted(files) == sorted([tmp_path / "h.md.txt", tmp_path / "y.md" / "i.txt"])
    # 取り込み元フォルダ自身のパスは比べない。
    assert len(ingest.list_source_files(tmp_path / "x", True, (".md",), exclude_keywords=[tmp_path.name])) == 1


def test_scan_keeps_only_paths_matching_include_keywords(tmp_path):
    for rel in ("a.md", "b.txt", "manual/c.md", "x/Manual/y/d.txt", "old/manual_e.md"):
        (tmp_path / rel).parent.mkdir(parents=True, exist_ok=True)
        (tmp_path / rel).write_text("x", encoding="utf-8")
    for workers in (1, 4):
        # どれかに当てはまるファイルだけ（フォルダ名でもよい。フォルダは絞らずにたどる）。除外が優先。
        files = [
            p for p, _ in ingest.scan_source_files(
                tmp_path, True, (".md", ".txt"), workers=workers, include_keywords=["MANUAL", "*.txt"], exclude_keywords=["/old/"]
            ).files
        ]
        assert files == [tmp_path / "b.txt", tmp_path / "manual" / "c.md", tmp_path / "x" / "Manual" / "y" / "d.txt"]
    files = ingest.list_source_files(tmp_path, True, (".md", ".txt"), include_keywords=["*.md"])
    assert sorted(files) == sorted([tmp_path / "a.md", tmp_path / "manual" / "c.md", tmp_path / "old" / "manual_e.md"])


def test_ingest_removes_files_matching_exclude_keywords(env):
    cfg, kwargs = env["cfg"], env["kwargs"]
    assert ingest.ingest_collection(cfg, "docs", **kwargs).added == 2
    with vcat.open_catalog(cfg.vector_db_dir) as catalog:
        # 空行と大文字小文字だけ違う重複は除く。
        assert catalog.update("docs", exclude_keywords=" weld\n\nWELD").exclude_keywords == ("weld",)
    # 除外キーワードに当てはまるようになった取り込み済みのファイルは消える。
    assert ingest.ingest_collection(cfg, "docs", **kwargs).removed == 1
    with vcat.open_catalog(cfg.vector_db_dir) as catalog:
        assert list(catalog.files("docs")) == [str(env["docs"] / "sub" / "room.txt")]
        catalog.update("docs", exclude_keywords=[])
    assert ingest.ingest_collection(cfg, "docs", **kwargs).added == 1
    # 対象キーワードに当てはまらなくなったファイルも消える。
    with vcat.open_catalog(cfg.vector_db_dir) as catalog:
        assert catalog.update("docs", include_keywords=["*.txt"]).include_keywords == ("*.txt",)
    assert ingest.ingest_collection(cfg, "docs", **kwargs).removed == 1
    with vcat.open_catalog(cfg.vector_db_dir) as catalog:
        assert list(catalog.files("docs")) == [str(env["docs"] / "sub" / "room.txt")]


def test_search_filters_by_collection_score_and_top_k(env, tmp_path):
    cfg, kwargs = env["cfg"], env["kwargs"]
    ingest.ingest_collection(cfg, "docs", **kwargs)
    with vcat.open_catalog(cfg.vector_db_dir) as catalog:
        catalog.create("empty", str(tmp_path))
    result = search.search(cfg, "会議室の予約", top_k=1, **kwargs)
    assert [h.source_path.endswith("room.txt") for h in result.hits] == [True]
    assert result.skipped == [("empty", "no documents ingested yet")]
    strict = make_config(cfg.vector_db_dir.parent, vector_db_backend=cfg.vector_db_backend, vector_db_min_score=0.99)
    assert [h.text for h in search.search(strict, "会議室の予約はポータルから行う。", **kwargs).hits] == ["会議室の予約はポータルから行う。"]
    with pytest.raises(search.SearchError):
        search.search(cfg, "x", collection="missing", **kwargs)
    with pytest.raises(search.SearchError):
        search.search(cfg, "  ", **kwargs)


def test_list_collections_without_catalog_does_not_create_it(tmp_path):
    cfg = make_config(tmp_path)
    assert search.list_collections(cfg) == []
    assert not cfg.vector_db_dir.exists()


def test_list_source_files_skips_hidden_files_and_folders(tmp_path):
    (tmp_path / "a.md").write_text("a", encoding="utf-8")
    (tmp_path / ".secret.md").write_text("x", encoding="utf-8")
    (tmp_path / ".git").mkdir()
    (tmp_path / ".git" / "b.md").write_text("x", encoding="utf-8")
    (tmp_path / "sub" / ".cache").mkdir(parents=True)
    (tmp_path / "sub" / ".cache" / "c.md").write_text("x", encoding="utf-8")
    (tmp_path / "sub" / "d.md").write_text("d", encoding="utf-8")
    files = ingest.list_source_files(tmp_path, True, (".md",))
    assert [p.relative_to(tmp_path).as_posix() for p in files] == ["a.md", "sub/d.md"]
